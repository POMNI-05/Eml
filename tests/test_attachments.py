import io
import unittest
from unittest.mock import patch

from eml_core import GmailMailbox, Selection
from eml_core.attachments import extract_attachment
from eml_core.documents import DocumentError, MAX_BYTES
from tests.fakes import FakeGmail, attach_file, raw_message


class AttachmentTests(unittest.TestCase):
    def test_external_and_inline_attachments_keep_their_exact_message_and_part(self):
        message = attach_file(raw_message("m1"), "notes.txt", b"External advice", attachment_id="a1")
        attach_file(message, "notes.txt", b"Inline notes", part_id="2")
        other = attach_file(raw_message("outside", labels=("L3",)), "private.txt", b"Other client")
        api = FakeGmail([message, other])
        api.attachment_data[("m1", "a1")] = b"External advice"
        result = GmailMailbox(api).select(Selection(label_ids=("L1",)))
        self.assertEqual(len(result.messages), 1)
        attachments = result.messages[0].attachment_contents
        self.assertEqual([a.part_id for a in attachments], ["1", "2"])
        self.assertEqual([a.parts[0].text for a in attachments], ["External advice", "Inline notes"])
        self.assertEqual(len(attachments[0].sha256), 64)
        calls = [params for name, params in api.calls if name == "attachments.get"]
        self.assertEqual(calls, [{"userId": "me", "messageId": "m1", "id": "a1"}])
        self.assertEqual(result.to_dict()["messages"][0]["attachment_contents"][1]["parts"][0]["location"], "text")

    def test_files_beyond_the_message_limit_or_subject_scope_are_not_downloaded(self):
        messages = [attach_file(raw_message(mid), "file.txt", attachment_id=mid) for mid in ["a", "b", "c"]]
        messages[0]["payload"]["headers"][2]["value"] = "Unrelated"
        api = FakeGmail(messages)
        api.attachment_data[("b", "b")] = b"Chosen content"
        result = GmailMailbox(api).select(Selection(label_ids=("L1",), subject_prefix="Subject", max_messages=1))
        self.assertEqual([m.message_id for m in result.messages], ["b"])
        self.assertEqual([params["messageId"] for name, params in api.calls if name == "attachments.get"], ["b"])

    def test_combined_attachment_limit_skips_further_downloads(self):
        messages = [attach_file(raw_message(mid), "file.txt", attachment_id=mid, size=10) for mid in ["a", "b"]]
        api = FakeGmail(messages)
        api.attachment_data[("a", "a")] = b"First file"
        api.attachment_data[("b", "b")] = b"Other file"
        with patch("eml_core.mailbox.MAX_FETCH_ATTACHMENT_BYTES", 15):
            result = GmailMailbox(api).select(Selection(label_ids=("L1",)))
        self.assertTrue(result.messages[0].attachment_contents[0].parts)
        self.assertFalse(result.messages[1].attachment_contents[0].parts)
        self.assertIn("50 MB", result.messages[1].attachment_contents[0].warnings[0])
        self.assertEqual(len([c for c in api.calls if c[0] == "attachments.get"]), 1)

    def test_unsupported_oversized_and_failed_downloads_are_explicit_without_losing_mail(self):
        message = attach_file(raw_message("m1"), "scan.png", b"image")
        attach_file(message, "large.pdf", attachment_id="too-large", size=MAX_BYTES + 1, part_id="2")
        attach_file(message, "missing.txt", attachment_id="missing", part_id="3")
        api = FakeGmail([message])
        result = GmailMailbox(api).select(Selection(message_ids=("m1",)))
        attachments = result.messages[0].attachment_contents
        self.assertEqual(len(attachments), 3)
        self.assertTrue(all(not a.parts and a.warnings for a in attachments))
        self.assertIn("OCR", attachments[0].warnings[0])
        self.assertIn("20 MB", attachments[1].warnings[0])
        self.assertEqual(len([c for c in api.calls if c[0] == "attachments.get"]), 1)
        self.assertEqual(result.messages[0].body, "Hello")

    def test_pdf_and_word_preserve_page_and_block_locations(self):
        from reportlab.pdfgen.canvas import Canvas
        from docx import Document
        pdf = io.BytesIO()
        canvas = Canvas(pdf)
        canvas.drawString(50, 750, "Advice was prepared for review.")
        canvas.showPage()
        canvas.showPage()
        canvas.save()
        parts, warnings = extract_attachment("advice.pdf", pdf.getvalue())
        self.assertEqual(parts[0].location, "page 1")
        self.assertIn("Advice was prepared", parts[0].text)
        self.assertTrue(any("Page 2" in warning for warning in warnings))
        document = Document()
        document.add_paragraph("Review findings")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Reconciliation completed"
        word = io.BytesIO()
        document.save(word)
        parts, _ = extract_attachment("report.docx", word.getvalue())
        self.assertEqual([p.location for p in parts], ["block 1", "block 2"])
        self.assertIn("Reconciliation completed", parts[1].text)

    def test_csv_and_excel_keep_rows_sheets_and_formula_text(self):
        from openpyxl import Workbook
        parts, _ = extract_attachment("records.csv", 'Name,Notes\nAlex,"Review, completed"'.encode("utf-16"))
        self.assertEqual(parts[1].location, "row 2")
        self.assertEqual(parts[1].text, "Alex | Review, completed")
        workbook = Workbook()
        workbook.active.title = "Review"
        workbook.active.append(["Activity", "Amount"])
        workbook.active.append(["Reconciliation", "=SUM(1,2)"])
        output = io.BytesIO()
        workbook.save(output)
        parts, warnings = extract_attachment("records.xlsx", output.getvalue())
        self.assertEqual(parts[1].location, "sheet Review, row 2")
        self.assertIn("=SUM(1,2)", parts[1].text)
        self.assertIn("not calculated", warnings[0])

    def test_image_only_pdf_is_not_reported_as_read(self):
        from pypdf import PdfWriter
        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf)
        message = attach_file(raw_message("m1"), "scan.pdf", pdf.getvalue())
        result = GmailMailbox(FakeGmail([message])).select(Selection(message_ids=("m1",)))
        attachment = result.messages[0].attachment_contents[0]
        self.assertFalse(attachment.parts)
        self.assertIn("OCR", attachment.warnings[0])
        with self.assertRaises(DocumentError):
            extract_attachment("empty.csv", b"\n\n")
