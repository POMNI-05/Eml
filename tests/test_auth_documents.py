import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from eml_core.auth import GmailAuth, CredentialError, validate_credentials
from eml_core.documents import BackgroundStore, DocumentError, extract_document, explicit_document_candidates
from eml_core.profiles import Candidate, Evidence, ProfileStore
from eml_core.search import build_client_search


def credential_bytes(client="test-client", secret="test-only"):
    return json.dumps({"installed": {
        "client_id": client + ".apps.googleusercontent.com",
        "client_secret": secret,
        "auth_uri": "https://accounts.google.com/o/oauth2/auth",
        "token_uri": "https://oauth2.googleapis.com/token",
        "redirect_uris": ["http://localhost"],
    }}).encode()


class CredentialTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.auth = GmailAuth(self.temp.name)

    def test_desktop_upload_is_saved_privately_without_echoing_secrets(self):
        self.auth.save_credentials(credential_bytes())
        self.assertEqual(os.stat(self.auth.credentials_path).st_mode & 0o777, 0o600)
        self.assertEqual(self.auth.status(), {
            "credentials_present": True, "authorization_present": False,
        })
        self.assertNotIn("client_secret", json.dumps(self.auth.status()))

    def test_invalid_or_wrong_oauth_types_do_not_create_files(self):
        for value in (b"not json", b'{"web":{}}', b'{"type":"service_account"}'):
            with self.subTest(value=value), self.assertRaises(CredentialError):
                self.auth.save_credentials(value)
        self.assertFalse(self.auth.credentials_path.exists())

    def test_uploaded_configuration_cannot_redirect_tokens_elsewhere(self):
        data = json.loads(credential_bytes())
        data["installed"]["token_uri"] = "https://outside.example/token"
        with self.assertRaises(CredentialError):
            validate_credentials(json.dumps(data).encode())

    def test_replacing_client_resets_token_but_same_configuration_keeps_it(self):
        self.auth.save_credentials(credential_bytes())
        self.auth.token_path.write_text('{"test":true}')
        self.assertFalse(self.auth.save_credentials(credential_bytes()))
        self.assertTrue(self.auth.token_path.exists())
        self.assertTrue(self.auth.save_credentials(credential_bytes(client="another")))
        self.assertFalse(self.auth.token_path.exists())

    def test_forgetting_sign_in_keeps_desktop_configuration(self):
        self.auth.save_credentials(credential_bytes())
        self.auth.token_path.write_text("{}")
        self.auth.disconnect()
        self.assertTrue(self.auth.credentials_path.exists())
        self.assertFalse(self.auth.token_path.exists())


class BackgroundTests(unittest.TestCase):
    def setUp(self):
        self.profiles = ProfileStore(":memory:")
        self.addCleanup(self.profiles.close)
        self.cid = self.profiles.create_client("Alex")
        self.background = BackgroundStore(self.profiles)

    def test_unicode_plaintext_and_explicit_details_reach_search_after_confirmation(self):
        text = "公司名称：Example Pty Ltd\n行业：批发\n邮箱：alex@client.example"
        doc = self.background.save(self.cid, "../client.txt", text.encode("utf-8-sig"))
        self.assertEqual(doc.filename, "client.txt")
        count = self.profiles.add_background_candidates(
            self.cid, [doc], explicit_document_candidates(doc)
        )
        self.assertEqual(count, 3)
        with self.assertRaises(ValueError):
            build_client_search(self.profiles.context(self.cid))
        for fact in self.profiles.list_facts(self.cid, "pending"):
            self.profiles.review(self.cid, fact["fact_id"], "confirmed")
        plan = build_client_search(self.profiles.context(self.cid))
        self.assertIn("Example Pty Ltd", plan.query)
        self.assertIn("from:", plan.query)
        self.assertIn("alex@client.example", plan.query)
        self.assertNotIn("{{", plan.query)
        self.assertTrue(plan.source_fact_ids)

    def test_utf16_text(self):
        parts, _ = extract_document("client.txt", "业务描述：跨境业务".encode("utf-16"))
        self.assertIn("跨境", parts[0].text)

    def test_search_terms_cannot_escape_their_quotes(self):
        plan = build_client_search(
            self.profiles.context(self.cid), ['Example\\" OR in:anywhere {other}']
        )
        self.assertNotIn("\\", plan.query)
        self.assertEqual(plan.query.count('"'), 2)
        self.assertEqual(plan.query.count("{"), 1)

    def test_docx_includes_tables_in_document_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Company: Example Pty Ltd")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Related company"
        table.cell(0, 1).text = "Another Pty Ltd"
        document.add_paragraph("Business: Wholesale")
        data = io.BytesIO()
        document.save(data)
        parts, _ = extract_document("background.docx", data.getvalue())
        self.assertEqual(len(parts), 3)
        self.assertIn("Related company | Another Pty Ltd", parts[1].text)
        self.assertTrue(parts[2].text.startswith("Business:"))

    def test_searchable_pdf_preserves_page_source_and_reports_unread_pages(self):
        from reportlab.pdfgen.canvas import Canvas
        data = io.BytesIO()
        canvas = Canvas(data)
        canvas.drawString(50, 750, "Company: Example Pty Ltd")
        canvas.showPage()
        canvas.showPage()
        canvas.save()
        parts, warnings = extract_document("client.pdf", data.getvalue())
        self.assertEqual(parts[0].location, "page 1")
        self.assertIn("Example Pty Ltd", parts[0].text)
        self.assertEqual(len(warnings), 1)

    def test_scanned_or_blank_pdf_does_not_become_a_successful_empty_document(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(DocumentError, "No searchable"):
            extract_document("scan.pdf", data.getvalue())

    def test_documents_and_originals_are_client_scoped(self):
        content = b"Company: Example Pty Ltd"
        first = self.background.save(self.cid, "client.txt", content)
        same = self.background.save(self.cid, "renamed.txt", content)
        self.assertEqual(first.document_id, same.document_id)
        other_id = self.profiles.create_client("Other")
        other = self.background.save(other_id, "client.txt", content)
        self.assertNotEqual(first.document_id, other.document_id)
        with self.assertRaises(DocumentError):
            self.background.get(other_id, first.document_id)
        self.assertEqual(self.background.original(self.cid, first.document_id), content)

    def test_invented_or_other_client_document_quotes_cannot_be_saved(self):
        doc = self.background.save(self.cid, "client.txt", b"Company: Example Pty Ltd")
        bad = Candidate("industry", "Retail", (Evidence(
            "background_document", "Invented retail quote",
            document_id=doc.document_id, location="text",
        ),))
        with self.assertRaises(ValueError):
            self.profiles.add_background_candidates(self.cid, [doc], [bad])
        other_id = self.profiles.create_client("Other")
        with self.assertRaises(ValueError):
            self.profiles.add_background_candidates(other_id, [doc], [])
        self.assertEqual(self.profiles.list_facts(self.cid), [])

    def test_unsupported_and_empty_documents_fail_clearly(self):
        for filename, content in [("file.xlsx", b"content"), ("empty.txt", b"")]:
            with self.subTest(filename=filename), self.assertRaises(DocumentError):
                extract_document(filename, content)


if __name__ == "__main__":
    unittest.main()
