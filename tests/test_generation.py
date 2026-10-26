import io
import json
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from eml_core.ai import OllamaClient, ModelError, BackgroundAnalysis, analyse_background, background_candidates
from eml_core import Email, ProfileStore, Selection, SelectionResult
from eml_core.documents import BackgroundStore
from invoice.models import BillingAnalysis
from invoice.service import DraftStore, generate_draft
from invoice.renderer import render_docx


class FakeModel:
    model = "test-local"
    context_size = 32768

    def __init__(self, bad_quote=False, wrong_id=False):
        self.payload = None
        self.bad_quote = bad_quote
        self.wrong_id = wrong_id

    def list_models(self):
        return [self.model]

    def generate(self, schema, system, payload):
        self.payload = payload
        if schema is BackgroundAnalysis:
            doc = payload["background_documents"][0]
            return schema.model_validate({
                "overview": "Alex represents Example Pty Ltd.",
                "facts": [
                    {"field": "entities", "value": {"name": "Example Pty Ltd"},
                     "evidence": [{"document_id": doc["document_id"], "location": "text",
                                   "quote": "Example Pty Ltd"}]},
                    {"field": "contacts", "value": {"name": "Alex", "email": "alex@client.example"},
                     "evidence": [{"document_id": doc["document_id"], "location": "text",
                                   "quote": "alex@client.example"}]},
                ],
                "questions": ["Check whether another related entity is in scope."],
            })
        message = payload["selected_email_evidence"][0]
        return schema.model_validate({
            "client_summary": "Example Pty Ltd is a wholesale client.",
            "activities": [{
                "entity": "Example Pty Ltd", "description": "Reviewed the bookkeeping records.",
                "status": "completed",
                "evidence": [{
                    "message_id": "outside" if self.wrong_id else message["message_id"],
                    "quote": "Invented completed work" if self.bad_quote else message["body"],
                }], "uncertainties": ["Check against prior billing records."],
            }],
            "open_questions": ["Confirm the time spent before billing."],
        })


class GenerationTests(unittest.TestCase):
    def setUp(self):
        self.store = ProfileStore(":memory:")
        self.addCleanup(self.store.close)
        self.cid = self.store.create_client("Alex")
        self.document = BackgroundStore(self.store).save(
            self.cid, "client.txt",
            b"Alex represents Example Pty Ltd. Contact alex@client.example. Wholesale business.",
        )
        self.message = Email(
            "m1", "t1", sender="staff@firm.example", to="alex@client.example",
            subject="Records review", date="2026-10-01",
            body="I completed the bookkeeping review.",
        )
        self.result = SelectionResult(
            "staff@firm.example", Selection(message_ids=("m1",)), [self.message]
        )

    def generate(self, client=None):
        return generate_draft(
            client or FakeModel(), self.store.context(self.cid), [self.document],
            self.result, ["m1"],
        )

    def test_background_and_exact_selected_mail_are_present_in_model_input(self):
        model = FakeModel()
        draft = self.generate(model)
        self.assertIn("Wholesale", model.payload["background_documents"][0]["parts"][0]["text"])
        self.assertEqual(
            [m["message_id"] for m in model.payload["selected_email_evidence"]], ["m1"]
        )
        self.assertFalse(draft["reviewed"])
        self.assertNotIn("hours", draft["analysis"]["activities"][0])

    def test_out_of_scope_or_invented_email_evidence_is_rejected(self):
        for model in (FakeModel(bad_quote=True), FakeModel(wrong_id=True)):
            with self.subTest(model=model), self.assertRaises(ModelError):
                self.generate(model)

    def test_background_quotes_turn_into_reviewable_candidates(self):
        analysis = analyse_background(FakeModel(), self.store.context(self.cid), [self.document])
        added = self.store.add_background_candidates(
            self.cid, [self.document], background_candidates(analysis)
        )
        self.assertEqual(added, 2)
        self.assertEqual(self.store.context(self.cid)["confirmed_facts"], [])

    def test_other_client_cannot_supply_documents_or_load_saved_draft(self):
        other = self.store.create_client("Other")
        with self.assertRaises(ModelError):
            generate_draft(FakeModel(), self.store.context(other), [self.document], self.result, ["m1"])
        drafts = DraftStore(self.store)
        draft = self.generate()
        drafts.save(draft)
        self.assertEqual(drafts.get(self.cid, draft["draft_id"])["input_digest"], draft["input_digest"])
        with self.assertRaises(ValueError):
            drafts.get(other, draft["draft_id"])

    def test_word_export_contains_edits_and_traceable_email_sources(self):
        from docx import Document
        draft = self.generate()
        edits = [{
            "Item": 1, "Include": True, "Entity": "Example Pty Ltd",
            "Description": "Reviewer corrected description",
            "Status": "in_progress", "Reviewer note": "Check remaining records.",
        }]
        document = Document(io.BytesIO(render_docx(draft, edits)))
        text = "\n".join(p.text for p in document.paragraphs)
        cells = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("draft", text)
        self.assertIn("m1", text)
        self.assertIn(self.message.body, text)
        self.assertIn("Reviewer corrected description", cells)
        self.assertIn("in progress", cells)

    def test_word_display_handles_control_characters_without_changing_saved_sources(self):
        from docx import Document
        self.message = Email("m1", "t1", body="Completed\x01 the review.")
        self.result.messages = [self.message]
        draft = self.generate()
        document = Document(io.BytesIO(render_docx(draft)))
        self.assertIn("Completed the review.", "\n".join(p.text for p in document.paragraphs))
        self.assertIn("\x01", draft["inputs"]["selected_email_evidence"][0]["body"])

    def test_schema_does_not_allow_the_model_to_add_fees(self):
        from pydantic import ValidationError
        value = self.generate()["analysis"]
        value["activities"][0]["hours"] = 5
        with self.assertRaises(ValidationError):
            BillingAnalysis.model_validate(value)

    def test_clear_request_wording_is_not_exported_as_established_completion(self):
        self.message = Email(
            "m1", "t1", body="Please prepare the October BAS after I send the records."
        )
        self.result.messages = [self.message]
        activity = self.generate()["analysis"]["activities"][0]
        self.assertEqual(activity["status"], "unclear")
        self.assertIn("completion is not established", activity["uncertainties"][-1])


class LocalHTTPTests(unittest.TestCase):
    def setUp(self):
        self.calls = []
        calls = self.calls

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_GET(self):
                result = {"models": [{"name": "test-local"}, {"name": "remote:cloud"}]}
                self.send_response(200)
                self.end_headers()
                self.wfile.write(json.dumps(result).encode())

            def do_POST(self):
                request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                calls.append(request)
                value = {"client_summary": "Test", "activities": [], "open_questions": []}
                self.send_response(200)
                self.end_headers()
                self.wfile.write(json.dumps({
                    "done": True, "message": {"content": json.dumps(value)}
                }).encode())

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)
        self.client = OllamaClient(
            "test-local", f"http://127.0.0.1:{self.server.server_port}"
        )

    def test_real_http_path_sends_schema_and_disables_streaming(self):
        self.assertEqual(self.client.list_models(), ["test-local"])
        response = self.client.generate(BillingAnalysis, "Use source evidence.", {"test": True})
        self.assertEqual(response.activities, [])
        self.assertFalse(self.calls[0]["stream"])
        self.assertEqual(self.calls[0]["format"]["title"], "BillingAnalysis")

    def test_oversized_inputs_are_not_silently_truncated_or_sent(self):
        with self.assertRaisesRegex(ModelError, "too large"):
            self.client.generate(BillingAnalysis, "System", {"large": "x" * 40_000})
        self.assertEqual(self.calls, [])

    def test_nonlocal_servers_and_cloud_models_are_not_selected(self):
        with self.assertRaises(ModelError):
            OllamaClient("model", "https://outside.example")
        with self.assertRaises(ModelError):
            OllamaClient("model:cloud")


if __name__ == "__main__":
    unittest.main()
