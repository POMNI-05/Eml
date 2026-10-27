"""Local attachment text extraction with per-file source locations."""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .documents import DocumentError, DocumentPart, MAX_BYTES, MAX_TEXT, extract_document

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"}


@dataclass(frozen=True)
class EmailAttachment:
    part_id: str
    filename: str
    mime_type: str
    size: int = 0
    sha256: str = ""
    parts: tuple[DocumentPart, ...] = ()
    warnings: tuple[str, ...] = ()


def extract_attachment(filename: str, content: bytes):
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        return extract_document(filename, content)
    if not content or len(content) > MAX_BYTES:
        raise DocumentError("Choose a non-empty attachment of at most 20 MB.")
    parts, length = [], 0

    def append(location, values):
        nonlocal length
        text = " | ".join("" if value is None else str(value) for value in values).strip()
        if not text.replace("|", "").strip():
            return
        length += len(text)
        if length > MAX_TEXT:
            raise DocumentError("Attachment text exceeds 200,000 characters. Choose a smaller file.")
        parts.append(DocumentPart(location, text))

    try:
        if suffix == ".csv":
            encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            for index, row in enumerate(csv.reader(io.StringIO(content.decode(encoding))), 1):
                append(f"row {index}", row)
        else:
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(p.file_size for p in archive.infolist()) > 60 * 1024 * 1024:
                    raise DocumentError("The expanded Excel workbook is too large.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                visited = 0
                for sheet in workbook:
                    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        visited += len(row)
                        if visited > 200_000:
                            raise DocumentError("The Excel workbook contains too many cells. Choose a smaller file.")
                        append(f"sheet {sheet.title}, row {index}", row)
            finally:
                workbook.close()
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("This attachment could not be read. Check its format and text encoding.") from exc
    if not parts:
        raise DocumentError("No readable rows found in this attachment.")
    warnings = ["Excel formulas are shown as formulas; they are not calculated. Charts and images were not read."] if suffix == ".xlsx" else []
    return parts, warnings
