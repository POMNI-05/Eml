"""Editable Word work records with separate activity and email evidence sections."""

import io
import re

from .validator import validate_review_edits


def _word_safe(value):
    """Preserve the saved source snapshot; remove invalid XML controls only in display."""
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)
    if isinstance(value, list):
        return [_word_safe(item) for item in value]
    if isinstance(value, dict):
        return {key: _word_safe(item) for key, item in value.items()}
    return value


def render_docx(draft: dict, edits: list[dict] | None = None,
                *, reviewed: bool = False, reviewer: str = "") -> bytes:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    draft, edits, reviewer = _word_safe(draft), _word_safe(edits), _word_safe(reviewer)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(1)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.title = "Billing work record"
    doc.core_properties.author = ""
    title = "Billing work record" if reviewed else "Billing work record draft"
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(draft["client_name"])
    scope = draft["inputs"]["selection_scope"]
    doc.add_paragraph(
        f"Evidence set: {scope['chosen_count']} selected messages from "
        f"{scope['loaded_count']} loaded messages. Mailbox: {scope['account_id']}."
    )
    chosen_scope = scope["selection"]
    start, end = chosen_scope.get("start_date"), chosen_scope.get("end_date")
    doc.add_paragraph(
        f"Email date filter: {start or 'open start'} through {end or 'open end'}."
        if start or end else "Email date filter: none."
    )
    if not scope["fetch_complete"]:
        doc.add_paragraph("The mailbox fetch was incomplete. This record covers only the chosen evidence.")
    doc.add_paragraph(
        f"Review status: {'reviewed by ' + (reviewer or 'reviewer') if reviewed else 'requires human review'}"
    )
    doc.add_paragraph("Time and fees are to be entered separately and are not inferred from correspondence.")

    analysis = draft["analysis"]
    doc.add_heading("Client background", 1)
    doc.add_paragraph(analysis["client_summary"] or "No background summary was supplied.")
    documents = draft["inputs"]["background_documents"]
    if documents:
        doc.add_paragraph("Background sources: " + "; ".join(d["filename"] for d in documents))
        for source in documents:
            for warning in source.get("warnings", []):
                doc.add_paragraph(f"{source['filename']}: {warning}")

    activities = analysis["activities"]
    if edits is None:
        edits = [
            {"Item": i + 1, "Include": True, "Entity": a["entity"],
             "Description": a["description"], "Status": a["status"], "Reviewer note": ""}
            for i, a in enumerate(activities)
        ]
    validate_review_edits(activities, edits)
    included = [row for row in edits if row.get("Include", True)]
    source_ids = []
    for row in included:
        index = int(row["Item"]) - 1
        for evidence in activities[index]["evidence"]:
            if evidence["message_id"] not in source_ids:
                source_ids.append(evidence["message_id"])
    refs = {mid: f"E{i + 1}" for i, mid in enumerate(source_ids)}

    doc.add_heading("Work activities", 1)
    if not included:
        doc.add_paragraph("No work activities have been included.")
    else:
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        widths = [3.45, 1.25, 1.05, 0.75]
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for cell, label, width in zip(
            table.rows[0].cells, ("Activity", "Entity", "Status", "Evidence"), widths
        ):
            cell.text = label
            cell.width = Inches(width)
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in included:
            activity = activities[int(row["Item"]) - 1]
            cells = table.add_row().cells
            description = str(row.get("Description") or "").strip()
            if row.get("Reviewer note"):
                description += "\nReviewer note: " + str(row["Reviewer note"])
            if activity["uncertainties"]:
                description += "\nTo verify: " + "; ".join(activity["uncertainties"])
            values = (
                description, str(row.get("Entity") or ""), str(row["Status"]).replace("_", " "),
                ", ".join(dict.fromkeys(refs[e["message_id"]] for e in activity["evidence"])),
            )
            for cell, value, width in zip(cells, values, widths):
                cell.text = value
                cell.width = Inches(width)
        for row_no, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "4")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "110")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "233D53" if row_no == 0 else ("F3F6F8" if row_no % 2 else "FFFFFF"))
                props.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.bold = row_no == 0
                        run.font.color.rgb = RGBColor(255, 255, 255) if row_no == 0 else RGBColor(0, 0, 0)

    if analysis["open_questions"]:
        doc.add_heading("Open questions", 1)
        for question in analysis["open_questions"]:
            doc.add_paragraph(question, "List Bullet")
    doc.add_heading("Email evidence", 1)
    messages = {m["message_id"]: m for m in draft["inputs"]["selected_email_evidence"]}
    for mid in source_ids:
        message = messages[mid]
        doc.add_heading(refs[mid], 2)
        doc.add_paragraph(f"Subject: {message['subject']}\nFrom: {message['sender']}\nDate: {message['date']}")
        doc.add_paragraph(f"Gmail message ID: {mid}")
        quotes = list(dict.fromkeys(
            e["quote"] for row in included
            for e in activities[int(row["Item"]) - 1]["evidence"] if e["message_id"] == mid
        ))
        for quote in quotes:
            doc.add_paragraph(quote)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
